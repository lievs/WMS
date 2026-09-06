"""Разбор файла плана отгрузок (выгрузка из общей таблицы MEVIAR).

Формат листов заранее не фиксирован жестко — заголовок с городами каждый
раз может съехать на строку/колонку, а сама таблица заливается заново раз
в 2 недели под новой датой в названии листа. Поэтому вместо фиксированных
номеров строк/колонок здесь идет поиск по смыслу: строка-заголовок находится
по ячейке "Баркод", дальше колонки-города определяются по тому, что их
подпись не похожа на служебную ("остатки", "отгружен / в пути" и т.п.).
"""

from openpyxl import load_workbook

# Подписи служебных/общих колонок (не города): и суб-колонки под городом
# (остаток/факт маркетплейса), и общие метаданные товара, которые могут
# встретиться правее штрихкода (GTIN, кол-во в коробе, план продаж и т.п.).
# Если подпись колонки не содержит ни одного из этих кусков — считаем ее
# названием нового города. Список специально с запасом (шире, чем нужно
# для текущего файла) — раз в 2 недели заливается новый файл, и лишняя
# устойчивость к формулировкам не помешает.
_SUBCOLUMN_MARKERS = (
    "остат", "отгруж", "путь", "факт", "план", "%", "gtin", "sku",
    "коробе", "производств", "хватает", "всего", "раскладк", "продаж",
    "дней", "поставк", "склад",
)

_SHEET_ALIASES = {
    "ozon": ("озон",),
    "wb": ("вб", "wb"),
}


def _norm(value):
    return str(value).strip() if value is not None else ""


def _find_plan_sheet(wb, marketplace):
    markers = _SHEET_ALIASES[marketplace]
    for name in wb.sheetnames:
        lower = name.lower()
        if "распределение" in lower and any(m in lower for m in markers):
            return name
    return None


def _find_header_row(ws, max_scan_rows=40):
    """Возвращает (row_idx, barcode_col) — строку и колонку с "Баркод"."""
    max_row = min(ws.max_row, max_scan_rows)
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = _norm(ws.cell(row=r, column=c).value).lower()
            if "баркод" in v:
                return r, c
    return None, None


def _find_label_col(ws, header_row, barcode_col, keyword):
    for c in range(1, barcode_col):
        v = _norm(ws.cell(row=header_row, column=c).value).lower()
        if keyword in v:
            return c
    return None


def _find_city_columns(ws, header_row, barcode_col):
    """[(col, city_name), ...] — только первая колонка каждой группы города
    (в ней лежит план по количеству), служебные колонки после нее пропускаем."""
    cities = []
    for c in range(barcode_col + 1, ws.max_column + 1):
        text = _norm(ws.cell(row=header_row, column=c).value)
        if not text:
            continue
        lower = text.lower()
        if any(marker in lower for marker in _SUBCOLUMN_MARKERS):
            continue
        cities.append((c, text))
    return cities


def _to_barcode_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_qty(value):
    if value is None or value == "":
        return None
    try:
        qty = float(value)
    except (TypeError, ValueError):
        return None
    return qty if qty > 0 else None


class ParsedPlan:
    def __init__(self, sheet_name):
        self.sheet_name = sheet_name
        self.cities = []  # list[str] в порядке появления
        self.rows = []  # list[dict]: barcode, article, size, city, qty


def parse_plan_sheet(file_stream, marketplace):
    """Возвращает ParsedPlan либо None, если подходящий лист не найден."""
    wb = load_workbook(file_stream, data_only=True)
    sheet_name = _find_plan_sheet(wb, marketplace)
    if not sheet_name:
        return None

    ws = wb[sheet_name]
    header_row, barcode_col = _find_header_row(ws)
    if header_row is None:
        return None

    article_col = _find_label_col(ws, header_row, barcode_col, "артикул")
    size_col = _find_label_col(ws, header_row, barcode_col, "размер")
    city_columns = _find_city_columns(ws, header_row, barcode_col)

    plan = ParsedPlan(sheet_name)
    plan.cities = [name for _, name in city_columns]

    for r in range(header_row + 1, ws.max_row + 1):
        barcode = _to_barcode_str(ws.cell(row=r, column=barcode_col).value)
        if not barcode:
            continue  # строки-подытоги ("ВСЕГО", "КАРДИГАНЫ" и т.п.) без штрихкода

        article = _norm(ws.cell(row=r, column=article_col).value) if article_col else ""
        size = _norm(ws.cell(row=r, column=size_col).value) if size_col else ""

        for col, city in city_columns:
            qty = _to_qty(ws.cell(row=r, column=col).value)
            if qty is None:
                continue
            plan.rows.append(
                {"barcode": barcode, "article": article, "size": size, "city": city, "qty": qty}
            )

    return plan
