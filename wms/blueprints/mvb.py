"""«МВБ Логистика» — отдельный раздел на базе WMS со своим входом (/mvb/login).

Клиент оформляет заявку на передачу коробов (забор транспортной компанией
или самопривоз) для отправки на СЦ Wildberries / Ozon. При оформлении каждый
короб получает собственный штрихкод, клиент печатает этикетки 58×40 и
клеит их на короба. Дальше короба сканируются поштучно: водитель при
заборе, склад МВБ при приемке и при отгрузке на СЦ, — и клиент видит у себя
статус каждого короба.

Пользователи с ролями MVB_ROLES видят только этот раздел; пользователи WMS
(кроме администраторов) — наоборот, сюда не попадают.
"""

import secrets
from datetime import date, datetime

from flask import (
    Blueprint, Response, abort, flash, jsonify, redirect, render_template, request, session,
    url_for,
)
from flask_login import current_user, login_user, logout_user

from ..extensions import db
from ..models import (
    MVB_BOX_STATUS_LABELS, MVB_BOX_STATUS_ORDER, MVB_BOX_STATUSES, MVB_DELIVERY_METHODS,
    MVB_MARKETPLACES, MVB_ROLES, MVB_TRIP_STATUSES, MvbBox, MvbBoxEvent, MvbClient, MvbOrder,
    MVB_PRICE_KINDS, MvbPallet, MvbPriceTier, MvbTrip, MvbTripStop, MvbVehicle, User,
)
from ..utils.http import content_disposition
from ..utils.labels_pdf import build_labels_batch_pdf
from ..utils.numbering import next_number
from ..utils.timezone import MOSCOW_OFFSET

bp = Blueprint("mvb", __name__)

# Эндпоинты, доступные без входа (проверяется в require_login приложения).
# Ссылка для наемного водителя на СЦ (без учетной записи) — по токену рейса.
MVB_PUBLIC_ENDPOINTS = {"mvb.login", "mvb.register", "mvb.trip_public", "mvb.trip_public_arrive", "mvb.trip_public_stop"}

MAX_BOXES_PER_ORDER = 500

# Режимы поштучного сканирования до склада: из каких статусов короб можно
# перевести в какой и каким ролям это разрешено. Администраторы могут всё.
# Дальше склада короба движутся в составе рейса (погрузка сканом, отправка
# и сдача на СЦ — см. раздел «Рейсы»).
SCAN_MODES = {
    "pickup": {
        "title": "Забор у клиента",
        "nav": "Скан забора",
        "to": "picked_up",
        "from": {"created"},
        "roles": {"mvb_driver", "mvb_staff", "mvb_admin"},
    },
    "receive": {
        "title": "Приемка на складе МВБ",
        "nav": "Приемка",
        "to": "received",
        "from": {"created", "picked_up", "not_delivered"},
        "roles": {"mvb_staff", "mvb_admin"},
    },
}

BOX_TIMESTAMP_FIELDS = {
    "picked_up": "picked_up_at",
    "received": "received_at",
    "loaded": "loaded_at",
    "shipped": "shipped_at",
    "delivered": "delivered_at",
}


# ---------- доступ ----------


def _is_client():
    return current_user.role == "mvb_client" and not current_user.is_admin


def _is_staff():
    return current_user.is_admin or current_user.role in ("mvb_staff", "mvb_admin")


def _is_driver():
    return current_user.role == "mvb_driver" and not current_user.is_admin


def _can_scan(mode):
    return current_user.is_admin or current_user.role in SCAN_MODES[mode]["roles"]


def _visible_orders_query():
    query = MvbOrder.query
    if _is_client():
        query = query.filter(MvbOrder.client_id == current_user.mvb_client_id)
    return query


def _get_order_or_404(order_id):
    order = _visible_orders_query().filter(MvbOrder.id == order_id).first()
    if order is None:
        abort(404)
    return order


@bp.before_request
def _restrict_to_mvb_users():
    if request.endpoint in MVB_PUBLIC_ENDPOINTS or not current_user.is_authenticated:
        return None
    if not (current_user.is_admin or current_user.role in MVB_ROLES):
        flash("Раздел «МВБ Логистика» вам не доступен", "danger")
        return redirect(url_for("main.index"))
    if _is_client() and current_user.mvb_client_id and _client_login_block(current_user):
        message = _client_login_block(current_user)
        logout_user()
        flash(message, "warning")
        return redirect(url_for("mvb.login"))
    if _is_client() and not current_user.mvb_client_id:
        logout_user()
        flash("Учетная запись не привязана к клиенту — обратитесь к администратору МВБ", "danger")
        return redirect(url_for("mvb.login"))
    return None


@bp.context_processor
def _inject():
    return {
        "MVB_BOX_STATUSES": MVB_BOX_STATUSES,
        "MVB_BOX_STATUS_LABELS": MVB_BOX_STATUS_LABELS,
        "MVB_MARKETPLACES": MVB_MARKETPLACES,
        "MVB_DELIVERY_METHODS": MVB_DELIVERY_METHODS,
        "MVB_ROLES": MVB_ROLES,
        "SCAN_MODES": SCAN_MODES,
        "mvb_is_client": current_user.is_authenticated and _is_client(),
        "mvb_can_scan": (lambda mode: current_user.is_authenticated and _can_scan(mode)),
        "mvb_is_staff": current_user.is_authenticated and _is_staff(),
        "mvb_pending_clients": (
            MvbClient.query.filter_by(approval="pending").count()
            if current_user.is_authenticated and _is_staff() else 0
        ),
        "MVB_TRIP_STATUSES": MVB_TRIP_STATUSES,
    }


# ---------- вход ----------


@bp.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated and (current_user.is_admin or current_user.role in MVB_ROLES):
        return redirect(url_for("mvb.index"))

    if request.method == "GET":
        return render_template("mvb/login.html")

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    user = User.query.filter_by(username=username).first()
    if (
        not user
        or not user.is_active_user
        or not user.check_password(password)
        or not (user.is_admin or user.role in MVB_ROLES)
    ):
        flash("Неверный логин или пароль", "danger")
        return render_template("mvb/login.html", username=username)

    blocked = _client_login_block(user)
    if blocked:
        flash(blocked, "warning")
        return render_template("mvb/login.html", username=username)

    login_user(user, remember=True)
    session["session_version"] = user.session_version or 0
    return redirect(url_for("mvb.index"))


def _client_login_block(user):
    """Почему клиент пока не может войти (регистрация не подтверждена,
    отклонена или клиент отключен); None — может."""
    if user.role != "mvb_client" or user.is_admin or user.mvb_client is None:
        return None
    client = user.mvb_client
    if client.approval == "pending":
        return "Регистрация на проверке у оператора МВБ — войти можно после подтверждения"
    if client.approval == "rejected":
        return "Регистрация отклонена — свяжитесь с МВБ Логистика"
    if not client.is_active:
        return "Учетная запись клиента отключена — свяжитесь с МВБ Логистика"
    return None


@bp.route("/register", methods=["GET", "POST"])
def register():
    """Самостоятельная регистрация клиента: компания + логин; клиент ждет
    подтверждения оператора (approval=pending), до этого войти нельзя."""
    if current_user.is_authenticated and (current_user.is_admin or current_user.role in MVB_ROLES):
        return redirect(url_for("mvb.index"))
    form = {k: request.form.get(k, "").strip() for k in (
        "name", "inn", "contact_name", "phone", "email", "address", "username",
    )}
    if request.method == "GET":
        return render_template("mvb/register.html", form=form)

    password = request.form.get("password", "")
    errors = []
    if not form["name"]:
        errors.append("Укажите название компании или ИП")
    inn = "".join(ch for ch in form["inn"] if ch.isdigit())
    if len(inn) not in (10, 12):
        errors.append("ИНН — 10 или 12 цифр")
    if not form["contact_name"]:
        errors.append("Укажите контактное лицо")
    if sum(ch.isdigit() for ch in form["phone"]) < 10:
        errors.append("Укажите телефон")
    if not form["username"]:
        errors.append("Придумайте логин")
    elif User.query.filter(db.func.lower(User.username) == form["username"].lower()).first():
        errors.append("Такой логин уже занят")
    if len(password) < 6:
        errors.append("Пароль — не короче 6 символов")
    elif password != request.form.get("password2", ""):
        errors.append("Пароли не совпадают")
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_template("mvb/register.html", form=form)

    client = MvbClient(
        name=form["name"], inn=inn, contact_name=form["contact_name"], phone=form["phone"],
        email=form["email"] or None, address=form["address"] or None, approval="pending",
    )
    user = User(username=form["username"], full_name=form["contact_name"], role="mvb_client", mvb_client=client)
    user.set_password(password)
    db.session.add_all([client, user])
    db.session.commit()
    flash("Заявка на регистрацию отправлена. Оператор МВБ проверит данные и подтвердит — после этого войдите со своим логином.", "success")
    return redirect(url_for("mvb.login"))


@bp.route("/logout", methods=["POST"])
def logout():
    logout_user()
    flash("Вы вышли из системы", "success")
    return redirect(url_for("mvb.login"))


@bp.route("/")
def index():
    if current_user.role == "mvb_driver" and not current_user.is_admin:
        return redirect(url_for("mvb.driver"))
    return redirect(url_for("mvb.orders"))


# ---------- заявки ----------


@bp.route("/orders")
def orders():
    query = _visible_orders_query()
    status = request.args.get("status", "")
    if status in ("draft", "confirmed", "cancelled"):
        query = query.filter(MvbOrder.status == status)
    client_id = request.args.get("client_id", type=int)
    if client_id and not _is_client():
        query = query.filter(MvbOrder.client_id == client_id)
    need_driver_filter = (
        MvbOrder.status == "confirmed", MvbOrder.delivery_method == "pickup", MvbOrder.driver_id.is_(None),
        MvbOrder.boxes.any(MvbBox.status == "created"),
    )
    if status == "need_driver":
        query = query.filter(*need_driver_filter)
    items = query.order_by(MvbOrder.created_at.desc()).limit(300).all()
    clients = [] if _is_client() else MvbClient.query.order_by(MvbClient.name).all()
    need_driver = 0 if _is_client() else MvbOrder.query.filter(*need_driver_filter).count()
    return render_template(
        "mvb/orders.html", orders=items, clients=clients, status=status, client_id=client_id,
        need_driver=need_driver, drivers=_active_drivers() if _is_staff() else [],
    )


def _parse_time(value):
    value = (value or "").strip()
    if not value:
        return None
    try:
        datetime.strptime(value, "%H:%M")
    except ValueError:
        return None
    return value


def _fill_order_from_form(order):
    """Заполняет заявку из формы; возвращает текст ошибки или None."""
    marketplace = request.form.get("marketplace", "wb")
    delivery_method = request.form.get("delivery_method", "pickup")
    if marketplace not in MVB_MARKETPLACES:
        return "Выберите маркетплейс"
    if delivery_method not in MVB_DELIVERY_METHODS:
        return "Выберите способ передачи коробов"
    try:
        box_count = int(request.form.get("box_count", ""))
    except ValueError:
        return "Укажите количество коробов"
    if box_count < 1 or box_count > MAX_BOXES_PER_ORDER:
        return f"Количество коробов — от 1 до {MAX_BOXES_PER_ORDER}"
    planned_date = None
    raw_date = request.form.get("planned_date", "").strip()
    if raw_date:
        try:
            planned_date = date.fromisoformat(raw_date)
        except ValueError:
            return "Некорректная дата"
    pickup_address = request.form.get("pickup_address", "").strip()
    if delivery_method == "pickup" and not pickup_address:
        return "Для забора укажите адрес"

    order.marketplace = marketplace
    order.delivery_method = delivery_method
    order.box_count = box_count
    order.destination = request.form.get("destination", "").strip() or None
    order.pickup_address = pickup_address or None
    order.planned_date = planned_date
    order.time_from = _parse_time(request.form.get("time_from"))
    order.time_to = _parse_time(request.form.get("time_to"))
    order.comment = request.form.get("comment", "").strip() or None
    return None


@bp.route("/orders/new", methods=["GET", "POST"])
def order_new():
    clients = [] if _is_client() else MvbClient.query.filter_by(is_active=True).order_by(MvbClient.name).all()
    if request.method == "GET":
        client = current_user.mvb_client if _is_client() else None
        return render_template(
            "mvb/order_form.html", order=None, clients=clients,
            default_address=(client.address if client else ""),
        )

    if _is_client():
        client_id = current_user.mvb_client_id
    else:
        client_id = request.form.get("client_id", type=int)
        if not client_id or not db.session.get(MvbClient, client_id):
            flash("Выберите клиента", "danger")
            return render_template("mvb/order_form.html", order=None, clients=clients, form=request.form)

    order = MvbOrder(client_id=client_id, created_by_id=current_user.id, status="draft")
    error = _fill_order_from_form(order)
    if error:
        flash(error, "danger")
        return render_template("mvb/order_form.html", order=None, clients=clients, form=request.form)
    order.number = next_number("mvb_order", "MVB-", 6)
    db.session.add(order)
    db.session.commit()
    flash(f"Заявка {order.number} создана. Проверьте и нажмите «Оформить» — коробам будут присвоены штрихкоды.", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>")
def order_detail(order_id):
    order = _get_order_or_404(order_id)
    drivers = _active_drivers() if _is_staff() else []
    return render_template(
        "mvb/order_detail.html", order=order, counts=order.status_counts(), drivers=drivers,
    )


@bp.route("/orders/<int:order_id>/edit", methods=["GET", "POST"])
def order_edit(order_id):
    order = _get_order_or_404(order_id)
    if order.status != "draft":
        flash("Изменить можно только черновик", "warning")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    if request.method == "GET":
        return render_template("mvb/order_form.html", order=order, clients=[])
    error = _fill_order_from_form(order)
    if error:
        db.session.rollback()
        flash(error, "danger")
        return render_template("mvb/order_form.html", order=order, clients=[], form=request.form)
    db.session.commit()
    flash("Заявка сохранена", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/confirm", methods=["POST"])
def order_confirm(order_id):
    """Оформление заявки: способ передачи фиксируется, каждому коробу
    присваивается собственный штрихкод «номер заявки-порядковый номер»."""
    order = _get_order_or_404(order_id)
    if order.status != "draft":
        flash("Заявка уже оформлена", "warning")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    for seq in range(1, order.box_count + 1):
        order.boxes.append(MvbBox(seq=seq, barcode=f"{order.number}-{seq:03d}", status="created"))
    order.status = "confirmed"
    order.confirmed_at = datetime.utcnow()
    _apply_prices(order)
    db.session.commit()
    flash(f"Заявка оформлена: присвоено штрихкодов — {order.box_count}. Распечатайте этикетки и наклейте на каждый короб.", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/cancel", methods=["POST"])
def order_cancel(order_id):
    order = _get_order_or_404(order_id)
    if order.status == "cancelled":
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    if any(box.status != "created" for box in order.boxes):
        flash("Нельзя отменить: часть коробов уже отсканирована", "danger")
        return redirect(url_for("mvb.order_detail", order_id=order.id))
    order.status = "cancelled"
    db.session.commit()
    flash("Заявка отменена", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/orders/<int:order_id>/labels.pdf")
def order_labels_pdf(order_id):
    """Этикетки 58×40 на все короба заявки (или на выбранные ?seq=1,2)."""
    order = _get_order_or_404(order_id)
    if order.status != "confirmed" or not order.boxes:
        abort(404)
    boxes = order.boxes
    seq_param = request.args.get("seq", "")
    if seq_param:
        try:
            wanted = {int(v) for v in seq_param.split(",") if v.strip()}
        except ValueError:
            abort(400)
        boxes = [b for b in boxes if b.seq in wanted]
    total = len(order.boxes)
    entries = [
        (
            box.barcode,
            f"{box.barcode}  ({box.seq}/{total})",
            f"{order.client.name[:28]} → {order.marketplace_label}",
        )
        for box in boxes
    ]
    pdf = build_labels_batch_pdf(entries, title_font_size=9, max_img_h_ratio=0.6)
    return Response(
        pdf,
        mimetype="application/pdf",
        headers={"Content-Disposition": content_disposition(f"{order.number}.pdf", "inline")},
    )


# ---------- сканирование ----------


# Заявка считается «новой» для водителя столько времени после оформления.
NEW_ORDER_HOURS = 3


def _driver_vehicle_and_load(user=None):
    """Авто водителя и сколько коробов сейчас у него в машине (забраны им,
    но еще не приняты на складе)."""
    user = user or current_user
    vehicle = (
        MvbVehicle.query.filter_by(driver_id=user.id, is_active=True)
        .order_by(MvbVehicle.id).first()
    )
    load = MvbBox.query.filter(MvbBox.status == "picked_up", MvbBox.picked_up_by_id == user.id).count()
    return vehicle, load


def _driver_active_trips(user=None):
    """Рейсы на СЦ, назначенные водителю и еще не завершенные."""
    user = user or current_user
    return (
        MvbTrip.query.filter(
            MvbTrip.driver_id == user.id,
            MvbTrip.status.in_(["assigned", "arrived", "loading", "departed"]),
        )
        .order_by(MvbTrip.planned_arrival_at)
        .all()
    )


@bp.route("/driver")
def driver():
    """Экран водителя на маршруте: сколько места в машине и лента заявок на
    забор (назначенные ему и свободные) с отметкой, влезает ли заявка;
    новые заявки подсвечиваются, страница сама обновляется."""
    candidates = (
        MvbOrder.query.filter(MvbOrder.status == "confirmed", MvbOrder.delivery_method == "pickup")
        .order_by(MvbOrder.confirmed_at.desc())
        .all()
    )
    vehicle, load = (None, 0)
    trips = []
    if _is_driver():
        vehicle, load = _driver_vehicle_and_load()
        trips = _driver_active_trips()
        # Водитель в рейсе на СЦ чужие заявки «по дороге» не берет — видит
        # только то, что оператор назначил ему самому.
        allowed = (current_user.id,) if trips else (None, current_user.id)
        candidates = [o for o in candidates if o.driver_id in allowed]
    free = (vehicle.capacity_boxes - load) if vehicle and vehicle.capacity_boxes else None
    now = datetime.utcnow()
    rows = []
    for order in candidates:
        remaining = sum(1 for b in order.boxes if b.status == "created")
        if not remaining:
            continue
        rows.append({
            "order": order,
            "remaining": remaining,
            "fits": None if free is None else remaining <= free,
            "is_new": bool(order.confirmed_at and (now - order.confirmed_at).total_seconds() < NEW_ORDER_HOURS * 3600),
            "mine": _is_driver() and order.driver_id == current_user.id,
        })
    # Сначала свои, затем новые, затем остальные по дате забора.
    rows.sort(key=lambda r: (not r["mine"], not r["is_new"], r["order"].planned_date or date.max))
    return render_template(
        "mvb/driver.html", rows=rows, trips=trips, on_trip=bool(trips), vehicle=vehicle, load=load, free=free,
        max_order_id=max((r["order"].id for r in rows), default=0),
    )


@bp.route("/driver/orders/<int:order_id>/take", methods=["POST"])
def driver_take(order_id):
    """Водитель берет свободную заявку на забор «по дороге»."""
    if not _is_driver():
        abort(403)
    order = MvbOrder.query.get_or_404(order_id)
    if order.status != "confirmed" or order.delivery_method != "pickup":
        abort(400)
    if order.driver_id not in (None, current_user.id):
        flash(f"Заявку {order.number} уже взял другой водитель", "warning")
        return redirect(url_for("mvb.driver"))
    if order.driver_id is None and _driver_active_trips():
        flash("Вы в рейсе на СЦ — заявки на забор назначает оператор", "warning")
        return redirect(url_for("mvb.driver"))
    vehicle, load = _driver_vehicle_and_load()
    remaining = sum(1 for b in order.boxes if b.status == "created")
    if vehicle and vehicle.capacity_boxes and remaining > vehicle.capacity_boxes - load:
        flash(f"Внимание: в машине свободно {max(vehicle.capacity_boxes - load, 0)} мест, а в заявке {remaining} кор.", "warning")
    order.driver_id = current_user.id
    db.session.commit()
    flash(f"Заявка {order.number} ваша — {order.pickup_address or 'адрес не указан'}", "success")
    return redirect(url_for("mvb.driver"))


@bp.route("/driver/orders/<int:order_id>/release", methods=["POST"])
def driver_release(order_id):
    if not _is_driver():
        abort(403)
    order = MvbOrder.query.get_or_404(order_id)
    if order.driver_id == current_user.id and not any(b.status != "created" for b in order.boxes):
        order.driver_id = None
        db.session.commit()
        flash(f"Вы отказались от заявки {order.number}", "success")
    return redirect(url_for("mvb.driver"))


@bp.route("/scan/<mode>")
def scan(mode):
    if mode not in SCAN_MODES:
        abort(404)
    if not _can_scan(mode):
        flash("Этот режим сканирования вам не доступен", "danger")
        return redirect(url_for("mvb.index"))
    return render_template("mvb/scan.html", mode=mode, mode_info=SCAN_MODES[mode])


@bp.route("/scan/<mode>", methods=["POST"])
def scan_box(mode):
    """Скан одного короба: JSON {ok, message, box...}. Повторный скан уже
    переведенного короба не ошибка — просто сообщаем, что он уже учтен."""
    if mode not in SCAN_MODES:
        abort(404)
    if not _can_scan(mode):
        return jsonify(ok=False, message="Нет доступа к этому режиму"), 403
    info = SCAN_MODES[mode]
    code = (request.form.get("barcode") or (request.get_json(silent=True) or {}).get("barcode") or "").strip()
    if not code:
        return jsonify(ok=False, message="Пустой штрихкод"), 400

    box = _find_box(code)
    imported_note = None
    if box is None and mode == "receive":
        # Свой короб WMS, перемещение которого еще не передано в МВБ, —
        # принимаем всё перемещение сразу по скану его этикетки.
        doc = _wms_movement_for_box_code(code)
        if doc is not None:
            order, errors = _import_movement(doc, "self")
            if order is not None:
                imported_note = f"Перемещение WMS {doc.number} принято в МВБ как заявка {order.number}"
                box = _find_box(code)
    if box is None:
        return jsonify(ok=False, message=f"Короб {code} не найден"), 404
    order = box.order
    payload = {
        "barcode": box.barcode,
        "order": order.number,
        "order_url": url_for("mvb.order_detail", order_id=order.id),
        "client": order.client.name,
        "seq": box.seq,
        "total": len(order.boxes),
    }
    if order.status != "confirmed":
        return jsonify(ok=False, message=f"Заявка {order.number} не оформлена или отменена", **payload), 409
    if mode == "pickup" and order.delivery_method != "pickup":
        return jsonify(ok=False, message="Это самопривоз — короб принимается на складе", **payload), 409

    if box.status == info["to"]:
        return jsonify(ok=True, already=True, message=f"Уже отмечен: {box.status_label}",
                       status=box.status_label, **payload)
    if box.status not in info["from"]:
        return jsonify(ok=False, message=f"Сейчас короб в статусе «{box.status_label}»", **payload), 409

    now = datetime.utcnow()
    returned = box.status == "not_delivered"
    _move_box(box, info["to"], now)
    if returned:
        # Вернулся с СЦ — снова доступен для погрузки в новый рейс.
        box.trip_id = box.trip_stop_id = box.pallet_id = None
        box.loaded_at = box.shipped_at = None
    load_note = None
    if mode == "pickup":
        box.picked_up_by_id = current_user.id
        if _is_driver() and order.driver_id is None:
            order.driver_id = current_user.id
    db.session.commit()
    if mode == "pickup" and _is_driver():
        vehicle, load = _driver_vehicle_and_load()
        if vehicle and vehicle.capacity_boxes:
            load_note = f"В машине {load} из {vehicle.capacity_boxes} кор."
    done = sum(
        1 for b in order.boxes
        if MVB_BOX_STATUS_ORDER.get(b.status, 0) >= MVB_BOX_STATUS_ORDER[info["to"]]
    )
    message = f"{box.barcode}: {box.status_label}"
    if imported_note:
        message = f"{imported_note}. {message}"
    return jsonify(
        ok=True, already=False, message=message, warning=load_note,
        status=box.status_label, done=done, **payload,
    )


# ---------- администрирование ----------


def _require_manage():
    if not current_user.can_manage_mvb():
        flash("Доступно только администратору МВБ", "danger")
        return False
    return True


@bp.route("/registrations")
def registrations():
    """Новые клиенты, зарегистрировавшиеся сами: оператор подтверждает или
    отклоняет."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    pending = MvbClient.query.filter_by(approval="pending").order_by(MvbClient.created_at).all()
    recent = (
        MvbClient.query.filter(MvbClient.approval.in_(["approved", "rejected"]), MvbClient.approved_at.isnot(None))
        .order_by(MvbClient.approved_at.desc()).limit(20).all()
    )
    return render_template("mvb/registrations.html", pending=pending, recent=recent)


@bp.route("/registrations/<int:client_id>/<action>", methods=["POST"])
def registration_action(client_id, action):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    client = db.session.get(MvbClient, client_id)
    if client is None or action not in ("approve", "reject"):
        abort(404)
    client.approval = "approved" if action == "approve" else "rejected"
    client.approved_at = datetime.utcnow()
    client.approved_by_id = current_user.id
    db.session.commit()
    if action == "approve":
        flash(f"Клиент «{client.name}» подтвержден — может входить и создавать заявки", "success")
    else:
        flash(f"Регистрация «{client.name}» отклонена", "warning")
    return redirect(url_for("mvb.registrations"))


@bp.route("/admin/clients", methods=["GET", "POST"])
def admin_clients():
    if not _require_manage():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        client_id = request.form.get("client_id", type=int)
        client = db.session.get(MvbClient, client_id) if client_id else MvbClient()
        if client is None:
            abort(404)
        name = request.form.get("name", "").strip()
        if not name:
            flash("Укажите название клиента", "danger")
            return redirect(url_for("mvb.admin_clients"))
        client.name = name
        client.inn = request.form.get("inn", "").strip() or None
        client.contact_name = request.form.get("contact_name", "").strip() or None
        client.phone = request.form.get("phone", "").strip() or None
        client.address = request.form.get("address", "").strip() or None
        if client_id:
            client.is_active = request.form.get("is_active") == "1"
        db.session.add(client)
        db.session.commit()
        flash(f"Клиент «{client.name}» сохранен", "success")
        return redirect(url_for("mvb.admin_clients"))
    clients = MvbClient.query.order_by(MvbClient.name).all()
    return render_template("mvb/admin_clients.html", clients=clients)


@bp.route("/admin/users", methods=["GET", "POST"])
def admin_users():
    if not _require_manage():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "")
        client_id = request.form.get("client_id", type=int)
        if not username or len(password) < 6:
            flash("Укажите логин и пароль не короче 6 символов", "danger")
        elif role not in MVB_ROLES:
            flash("Выберите роль", "danger")
        elif role == "mvb_client" and not (client_id and db.session.get(MvbClient, client_id)):
            flash("Для роли «Клиент» выберите клиента", "danger")
        elif User.query.filter_by(username=username).first():
            flash("Такой логин уже занят", "danger")
        else:
            user = User(
                username=username,
                full_name=request.form.get("full_name", "").strip() or None,
                role=role,
                is_admin=False,
                mvb_client_id=client_id if role == "mvb_client" else None,
                allowed_sections="none",
            )
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash(f"Пользователь «{username}» создан", "success")
        return redirect(url_for("mvb.admin_users"))
    users = User.query.filter(User.role.in_(list(MVB_ROLES))).order_by(User.username).all()
    clients = MvbClient.query.filter_by(is_active=True).order_by(MvbClient.name).all()
    return render_template("mvb/admin_users.html", users=users, clients=clients)


def _get_mvb_user_or_404(user_id):
    user = db.session.get(User, user_id)
    if user is None or user.role not in MVB_ROLES:
        abort(404)
    return user


@bp.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
def admin_user_toggle(user_id):
    if not _require_manage():
        return redirect(url_for("mvb.index"))
    user = _get_mvb_user_or_404(user_id)
    if user.id == current_user.id:
        flash("Нельзя отключить самого себя", "danger")
        return redirect(url_for("mvb.admin_users"))
    user.is_active_user = not user.is_active_user
    if not user.is_active_user:
        user.session_version = (user.session_version or 0) + 1
    db.session.commit()
    flash(f"Пользователь «{user.username}» {'включен' if user.is_active_user else 'отключен'}", "success")
    return redirect(url_for("mvb.admin_users"))


@bp.route("/admin/users/<int:user_id>/password", methods=["POST"])
def admin_user_password(user_id):
    if not _require_manage():
        return redirect(url_for("mvb.index"))
    user = _get_mvb_user_or_404(user_id)
    password = request.form.get("password", "")
    if len(password) < 6:
        flash("Пароль не короче 6 символов", "danger")
        return redirect(url_for("mvb.admin_users"))
    user.set_password(password)
    user.session_version = (user.session_version or 0) + 1
    db.session.commit()
    flash(f"Пароль для «{user.username}» изменен", "success")
    return redirect(url_for("mvb.admin_users"))


# ---------- общие помощники этапа 2 ----------


def _require_staff():
    if not _is_staff():
        flash("Доступно только складу МВБ", "danger")
        return False
    return True


def _active_drivers():
    return (
        User.query.filter(User.role == "mvb_driver", User.is_active_user.is_(True))
        .order_by(User.full_name, User.username)
        .all()
    )


def _move_box(box, status, now):
    box.status = status
    if status in BOX_TIMESTAMP_FIELDS:
        setattr(box, BOX_TIMESTAMP_FIELDS[status], now)
    user_id = current_user.id if current_user.is_authenticated else None
    db.session.add(MvbBoxEvent(box=box, status=status, user_id=user_id, created_at=now))


def _parse_dt(value):
    value = (value or "").strip()
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def _local_to_utc(value):
    """Время из формы вводится по Москве (как отображается везде в WMS), а
    хранится в UTC — как и остальные отметки времени."""
    dt = _parse_dt(value)
    return dt - MOSCOW_OFFSET if dt else None


def _same_direction(box_order, marketplace, destination):
    return box_order.marketplace == marketplace and (
        (box_order.destination or "").strip().lower() == (destination or "").strip().lower()
    )


# ---------- назначение водителя ----------


@bp.route("/orders/<int:order_id>/driver", methods=["POST"])
def order_assign_driver(order_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    order = _get_order_or_404(order_id)
    driver_id = request.form.get("driver_id", type=int)
    if driver_id:
        driver = db.session.get(User, driver_id)
        if driver is None or driver.role != "mvb_driver":
            abort(400)
        order.driver_id = driver.id
        flash(f"На забор назначен водитель {driver.display_name()}", "success")
    else:
        order.driver_id = None
        flash("Водитель снят с заявки", "success")
    db.session.commit()
    if request.form.get("back") == "orders":
        return redirect(url_for("mvb.orders", status=request.form.get("status", "")))
    return redirect(url_for("mvb.order_detail", order_id=order.id))


# ---------- транспорт ----------


@bp.route("/vehicles", methods=["GET", "POST"])
def vehicles():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        vehicle_id = request.form.get("vehicle_id", type=int)
        vehicle = db.session.get(MvbVehicle, vehicle_id) if vehicle_id else MvbVehicle()
        if vehicle is None:
            abort(404)
        plate = request.form.get("plate", "").strip().upper()
        try:
            capacity = int(request.form.get("capacity_boxes") or 0)
        except ValueError:
            capacity = -1
        if not plate or capacity < 0:
            flash("Укажите госномер и вместимость (число коробов)", "danger")
            return redirect(url_for("mvb.vehicles"))
        driver_id = request.form.get("driver_id", type=int)
        vehicle.plate = plate
        vehicle.model = request.form.get("model", "").strip() or None
        vehicle.carrier = request.form.get("carrier", "").strip() or None
        vehicle.capacity_boxes = capacity
        vehicle.driver_id = driver_id or None
        if vehicle_id:
            vehicle.is_active = request.form.get("is_active") == "1"
        db.session.add(vehicle)
        db.session.commit()
        flash(f"Транспорт {vehicle.plate} сохранен", "success")
        return redirect(url_for("mvb.vehicles"))
    return render_template(
        "mvb/vehicles.html",
        vehicles=MvbVehicle.query.order_by(MvbVehicle.is_active.desc(), MvbVehicle.plate).all(),
        drivers=_active_drivers(),
    )


# ---------- паллеты ----------


@bp.route("/pallets", methods=["GET", "POST"])
def pallets():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        marketplace = request.form.get("marketplace", "")
        if marketplace not in MVB_MARKETPLACES:
            flash("Выберите маркетплейс", "danger")
            return redirect(url_for("mvb.pallets"))
        pallet = MvbPallet(
            number=next_number("mvb_pallet", "PLT-", 6),
            marketplace=marketplace,
            destination=request.form.get("destination", "").strip() or None,
            created_by_id=current_user.id,
        )
        db.session.add(pallet)
        db.session.commit()
        return redirect(url_for("mvb.pallet_detail", pallet_id=pallet.id))
    items = MvbPallet.query.order_by(MvbPallet.created_at.desc()).limit(200).all()
    return render_template("mvb/pallets.html", pallets=items)


@bp.route("/pallets/<int:pallet_id>")
def pallet_detail(pallet_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    return render_template("mvb/pallet_detail.html", pallet=pallet)


@bp.route("/pallets/<int:pallet_id>/scan", methods=["POST"])
def pallet_scan(pallet_id):
    """Скан короба на паллету: только принятые на складе короба того же
    направления, еще не погруженные; с другой паллеты короб переносится."""
    if not _is_staff():
        return jsonify(ok=False, message="Нет доступа"), 403
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    code = (request.form.get("barcode") or "").strip()
    box = MvbBox.query.filter(db.func.upper(MvbBox.barcode) == code.upper()).first()
    if box is None:
        return jsonify(ok=False, message=f"Короб {code} не найден"), 404
    if box.status != "received":
        return jsonify(ok=False, message=f"Короб в статусе «{box.status_label}» — на паллету только принятые на складе"), 409
    if not _same_direction(box.order, pallet.marketplace, pallet.destination):
        return jsonify(ok=False, message=f"Другое направление: {box.order.marketplace_label} · {box.order.destination or '—'}"), 409
    if box.pallet_id == pallet.id:
        return jsonify(ok=True, already=True, message=f"{box.barcode} уже на этой паллете", count=len(pallet.boxes))
    moved_from = box.pallet.number if box.pallet else None
    box.pallet = pallet
    db.session.commit()
    message = f"{box.barcode} → {pallet.number}" + (f" (снят с {moved_from})" if moved_from else "")
    return jsonify(ok=True, already=False, message=message, count=len(pallet.boxes))


@bp.route("/pallets/<int:pallet_id>/remove/<int:box_id>", methods=["POST"])
def pallet_remove_box(pallet_id, box_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    box = db.session.get(MvbBox, box_id)
    if box is None or box.pallet_id != pallet_id or box.status != "received":
        abort(400)
    box.pallet_id = None
    db.session.commit()
    return redirect(url_for("mvb.pallet_detail", pallet_id=pallet_id))


@bp.route("/pallets/<int:pallet_id>/label.pdf")
def pallet_label_pdf(pallet_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    pallet = db.session.get(MvbPallet, pallet_id) or abort(404)
    subtitle = f"{pallet.marketplace_label} · {pallet.destination or ''} · {len(pallet.boxes)} кор."
    pdf = build_labels_batch_pdf([(pallet.number, pallet.number, subtitle)], title_font_size=12, max_img_h_ratio=0.6)
    return Response(
        pdf, mimetype="application/pdf",
        headers={"Content-Disposition": content_disposition(f"{pallet.number}.pdf", "inline")},
    )


# ---------- отправка на СЦ: готово к отправке и рейсы ----------


def _ready_boxes_query():
    return (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.status == "received", MvbBox.trip_id.is_(None))
        .order_by(MvbBox.received_at)
    )


def _direction_key(marketplace, destination):
    return f"{marketplace}|{(destination or '').strip()}"


def _ready_groups():
    """Принятые и еще не погруженные короба по направлениям (маркетплейс +
    СЦ), самые давние сверху (FIFO)."""
    groups = {}
    for box in _ready_boxes_query().all():
        key = _direction_key(box.order.marketplace, box.order.destination)
        group = groups.setdefault(key, {
            "key": key, "marketplace": box.order.marketplace,
            "destination": (box.order.destination or "").strip(), "boxes": 0,
            "oldest": box.received_at, "pallets": set(), "clients": set(),
        })
        group["boxes"] += 1
        group["clients"].add(box.order.client.name)
        if box.pallet_id:
            group["pallets"].add(box.pallet.number)
    return sorted(groups.values(), key=lambda g: g["oldest"])


def _ready_order_items(selected=None):
    """Готовые к отправке короба по заявкам — заявка внутри направления
    едет целиком, поэтому компонуем рейсы заявками. Порядок FIFO: сначала
    самые давние направления, внутри — самые давние заявки."""
    items = {}
    for box in _ready_boxes_query().all():
        key = _direction_key(box.order.marketplace, box.order.destination)
        if selected is not None and key not in selected:
            continue
        item = items.setdefault(box.order_id, {
            "order": box.order, "key": key, "marketplace": box.order.marketplace,
            "destination": (box.order.destination or "").strip(), "boxes": 0, "oldest": box.received_at,
        })
        item["boxes"] += 1
    group_oldest = {}
    for item in items.values():
        group_oldest[item["key"]] = min(group_oldest.get(item["key"], item["oldest"]), item["oldest"])
    return sorted(items.values(), key=lambda i: (group_oldest[i["key"]], i["key"], i["oldest"]))


def _pack_orders(items, capacity):
    """Компоновка рейсов без разбиения заявок: каждая заявка целиком идет в
    первую машину, где хватает места, иначе — в новую. Заявка больше
    вместимости едет отдельной машиной (сверх вместимости)."""
    bins = []
    for item in items:
        target = next((b for b in bins if b["boxes"] + item["boxes"] <= capacity), None)
        if target is None:
            target = {"items": [], "boxes": 0}
            bins.append(target)
        target["items"].append(item)
        target["boxes"] += item["boxes"]
    return bins


def _vehicle_capacities():
    return sorted({
        v.capacity_boxes for v in MvbVehicle.query.filter_by(is_active=True).all() if v.capacity_boxes
    })


@bp.route("/dispatch")
def dispatch():
    """Готово к отправке: по каждому направлению — сколько коробов ждет и
    сколько машин выбранной вместимости нужно; отмеченные направления можно
    отправить одним рейсом-маршрутом или разбить на рейсы по наполненности
    авто."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    rows = _ready_groups()
    capacities = _vehicle_capacities()
    capacity = request.args.get("capacity", type=int) or (capacities[-1] if capacities else 0)
    open_trips = {}
    for trip in MvbTrip.query.filter(MvbTrip.status.in_(["searching", "assigned", "arrived", "loading"])).all():
        for stop in trip.stops:
            open_trips.setdefault(_direction_key(stop.marketplace, stop.destination), []).append(trip)
    items = _ready_order_items()
    for row in rows:
        row["trips"] = open_trips.get(row["key"], [])
        row["orders"] = [i for i in items if i["key"] == row["key"]]
        row["vehicles"] = len(_pack_orders(row["orders"], capacity)) if capacity else None
    total = sum(r["boxes"] for r in rows)
    proposal = _pack_orders(items, capacity) if capacity else []
    return render_template(
        "mvb/dispatch.html", rows=rows, capacity=capacity, capacities=capacities, total=total,
        total_vehicles=len(proposal) if capacity else None, proposal=proposal,
    )


@bp.route("/trips")
def trips():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    status = request.args.get("status", "active")
    query = MvbTrip.query
    if status == "active":
        query = query.filter(MvbTrip.status.notin_(["delivered", "cancelled"]))
    elif status in MVB_TRIP_STATUSES:
        query = query.filter(MvbTrip.status == status)
    items = query.order_by(MvbTrip.created_at.desc()).limit(300).all()
    counts = dict(
        db.session.query(MvbTrip.status, db.func.count(MvbTrip.id)).group_by(MvbTrip.status).all()
    )
    return render_template("mvb/trips.html", trips=items, status=status, counts=counts)


def _add_stop(trip, marketplace, destination):
    stop = trip.stop_for(marketplace, destination)
    if stop is None:
        stop = MvbTripStop(
            marketplace=marketplace, destination=(destination or "").strip() or None,
            planned_boxes=0, seq=(max((s.seq for s in trip.stops), default=0) + 1),
        )
        trip.stops.append(stop)
    return stop


@bp.route("/trips/new", methods=["POST"])
def trip_new():
    """Рейсы по отмеченным направлениям (в порядке FIFO).

    mode=single — один рейс-маршрут по всем точкам; mode=fill — компоновка
    по наполненности авто целыми заявками (заявка не делится между машинами,
    см. _pack_orders). Каждый рейс создается в статусе «Поиск авто», заявки
    запоминаются как запланированные в него (planned_trip)."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    selected = set(request.form.getlist("dir"))
    groups = [g for g in _ready_groups() if g["key"] in selected]
    # Направление, для которого коробов уже нет (успели погрузить), все
    # равно можно добавить точкой — без плана.
    known = {g["key"] for g in groups}
    for key in selected - known:
        marketplace, _, destination = key.partition("|")
        if marketplace in MVB_MARKETPLACES:
            groups.append({"key": key, "marketplace": marketplace, "destination": destination, "boxes": 0})
    if not groups:
        flash("Отметьте хотя бы одно направление", "warning")
        return redirect(url_for("mvb.dispatch"))

    mode = request.form.get("mode", "single")
    capacity = request.form.get("capacity", type=int) or 0
    created = []

    def new_trip():
        trip = MvbTrip(
            number=next_number("mvb_trip", "RS-", 6), planned_boxes=0, status="searching",
            created_by_id=current_user.id,
        )
        db.session.add(trip)
        created.append(trip)
        return trip

    items = _ready_order_items(selected)
    oversized = []

    def plan(trip, item):
        stop = _add_stop(trip, item["marketplace"], item["destination"])
        stop.planned_boxes += item["boxes"]
        trip.planned_boxes += item["boxes"]
        item["order"].planned_trip = trip

    if mode == "fill" and capacity > 0:
        for bin_ in _pack_orders(items, capacity):
            trip = new_trip()
            for item in bin_["items"]:
                plan(trip, item)
                if item["boxes"] > capacity:
                    oversized.append(item["order"].number)
    else:
        trip = new_trip()
        for item in items:
            plan(trip, item)
    # Направления без коробов — точкой в первый рейс, без плана.
    first = created[0] if created else new_trip()
    for group in groups:
        if group["boxes"] == 0:
            _add_stop(first, group["marketplace"], group["destination"])
    db.session.commit()
    if oversized:
        flash(f"Заявки больше вместимости машины едут отдельной машиной целиком: {', '.join(oversized)}", "warning")
    if len(created) == 1:
        flash(f"Рейс {created[0].number} создан: {created[0].route_label()} — статус «Поиск авто»", "success")
        return redirect(url_for("mvb.trip_detail", trip_id=created[0].id))
    flash(
        f"Создано рейсов: {len(created)} ({', '.join(t.number for t in created)}) по {capacity} кор., "
        "заявки не разбиты — все в статусе «Поиск авто»", "success",
    )
    return redirect(url_for("mvb.trips", status="searching"))


def _get_trip_or_404(trip_id):
    trip = db.session.get(MvbTrip, trip_id)
    if trip is None:
        abort(404)
    if not _is_staff() and not (_is_driver() and trip.driver_id == current_user.id):
        abort(404)
    return trip


@bp.route("/trips/<int:trip_id>")
def trip_detail(trip_id):
    trip = _get_trip_or_404(trip_id)
    ready = {}
    if _is_staff():
        groups = {g["key"]: g["boxes"] for g in _ready_groups()}
        ready = {stop.id: groups.get(_direction_key(stop.marketplace, stop.destination), 0) for stop in trip.stops}
    if _is_staff() and not trip.access_token:
        trip.access_token = secrets.token_urlsafe(16)
        db.session.commit()
    return render_template(
        "mvb/trip_detail.html", trip=trip, ready=ready,
        driver_link=url_for("mvb.trip_public", token=trip.access_token, _external=True) if _is_staff() else None,
        vehicles=MvbVehicle.query.filter_by(is_active=True).order_by(MvbVehicle.plate).all() if _is_staff() else [],
        drivers=_active_drivers() if _is_staff() else [],
    )


TRIP_EDITABLE_STATUSES = ("searching", "assigned", "arrived", "loading")


@bp.route("/trips/<int:trip_id>/stops", methods=["POST"])
def trip_add_stop(trip_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    trip = _get_trip_or_404(trip_id)
    marketplace = request.form.get("marketplace", "")
    if trip.status not in TRIP_EDITABLE_STATUSES or marketplace not in MVB_MARKETPLACES:
        abort(400)
    stop = _add_stop(trip, marketplace, request.form.get("destination", ""))
    db.session.commit()
    flash(f"Точка «{stop.label()}» в маршруте", "success")
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


def _apply_stop_result(trip, stop, action, comment):
    """Итог на точке: «Сдано на СЦ» (deliver) или «Не сдано» с причиной
    (reject — короба возвращаются на склад МВБ). Возвращает (категория,
    сообщение) для flash."""
    if trip.status != "departed" or stop.result:
        return "warning", "Эта точка уже отмечена или рейс еще не в пути"
    comment = (comment or "").strip()
    if action == "reject" and not comment:
        return "danger", "Укажите причину, почему не сдано"
    now = datetime.utcnow()
    stop.result = "delivered" if action == "deliver" else "rejected"
    stop.delivered_at = now
    stop.delivered_by_id = current_user.id if current_user.is_authenticated else None
    stop.delivery_comment = comment or None
    for box in stop.boxes:
        if box.status == "shipped":
            _move_box(box, "delivered" if action == "deliver" else "not_delivered", now)
    if all(s.result for s in trip.stops):
        trip.status = "delivered"
        trip.delivered_at = now
    db.session.commit()
    if action == "deliver":
        return "success", f"Сдано: {stop.label()} — {len(stop.boxes)} кор."
    return "warning", f"Не сдано: {stop.label()} — {len(stop.boxes)} кор. везите обратно на склад МВБ"


@bp.route("/trips/<int:trip_id>/stops/<int:stop_id>/<action>", methods=["POST"])
def trip_stop_action(trip_id, stop_id, action):
    """Порядок точек (up/down), удаление пустой точки (remove) и итог
    водителя на точке: «Сдано на СЦ» (deliver) или «Не сдано» с причиной
    (reject — короба возвращаются на склад МВБ)."""
    trip = _get_trip_or_404(trip_id)
    stop = db.session.get(MvbTripStop, stop_id)
    if stop is None or stop.trip_id != trip.id:
        abort(404)
    now = datetime.utcnow()

    if action in ("deliver", "reject"):
        category, message = _apply_stop_result(trip, stop, action, request.form.get("comment", ""))
        flash(message, category)
        if _is_driver():
            return redirect(url_for("mvb.driver"))
        return redirect(url_for("mvb.trip_detail", trip_id=trip.id))

    if not _is_staff():
        abort(403)
    if trip.status not in TRIP_EDITABLE_STATUSES:
        abort(400)
    stops = list(trip.stops)
    index = stops.index(stop)
    if action in ("up", "down"):
        other = index - 1 if action == "up" else index + 1
        if 0 <= other < len(stops):
            stops[index].seq, stops[other].seq = stops[other].seq, stops[index].seq
    elif action == "remove":
        if stop.boxes:
            flash("На эту точку уже погружены короба — сначала отмените рейс", "danger")
            return redirect(url_for("mvb.trip_detail", trip_id=trip.id))
        db.session.delete(stop)
    else:
        abort(404)
    db.session.commit()
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


@bp.route("/trips/<int:trip_id>/plan", methods=["POST"])
def trip_plan(trip_id):
    """Авто найдено: транспорт, водитель и плановое время подачи/погрузки."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    trip = _get_trip_or_404(trip_id)
    if trip.status not in TRIP_EDITABLE_STATUSES:
        abort(400)
    vehicle_id = request.form.get("vehicle_id", type=int)
    vehicle = db.session.get(MvbVehicle, vehicle_id) if vehicle_id else None
    driver_id = request.form.get("driver_id", type=int)
    trip.vehicle = vehicle
    driver = db.session.get(User, driver_id) if driver_id else (vehicle.driver if vehicle else None)
    if driver_id and (driver is None or driver.role != "mvb_driver"):
        abort(400)
    trip.driver = driver
    trip.planned_arrival_at = _local_to_utc(request.form.get("planned_arrival_at"))
    trip.planned_load_start_at = _local_to_utc(request.form.get("planned_load_start_at"))
    trip.planned_load_end_at = _local_to_utc(request.form.get("planned_load_end_at"))
    trip.comment = request.form.get("comment", "").strip() or None
    # Наемный водитель (обычно случайный) — без учетной записи.
    trip.driver_name = request.form.get("driver_name", "").strip() or None
    trip.driver_phone = request.form.get("driver_phone", "").strip() or None
    trip.car_plate = request.form.get("car_plate", "").strip().upper() or None
    trip.capacity_boxes = request.form.get("capacity_boxes", type=int) or None
    if not trip.access_token:
        trip.access_token = secrets.token_urlsafe(16)
    if trip.has_transport() and trip.status == "searching":
        trip.status = "assigned"
    if not trip.has_transport() and trip.status == "assigned":
        trip.status = "searching"
    db.session.commit()
    flash("План рейса сохранен", "success")
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


@bp.route("/trips/<int:trip_id>/<action>", methods=["POST"])
def trip_action(trip_id, action):
    """Фактические отметки рейса. Отправка переводит погруженные короба в
    «В пути на СЦ» (пустые точки убираются из маршрута); водитель рейса
    может отметить подачу авто. Сдача — по точкам (trip_stop_action)."""
    trip = _get_trip_or_404(trip_id)
    now = datetime.utcnow()
    if not _is_staff() and action != "arrive":
        abort(403)

    if action == "arrive" and trip.status == "assigned":
        trip.status = "arrived"
        trip.arrived_at = now
    elif action == "start_loading" and trip.status in ("assigned", "arrived"):
        trip.status = "loading"
        trip.arrived_at = trip.arrived_at or now
        trip.load_started_at = now
    elif action == "depart" and trip.status in ("assigned", "arrived", "loading"):
        if not trip.boxes:
            flash("В рейсе нет погруженных коробов", "danger")
            return redirect(url_for("mvb.trip_detail", trip_id=trip.id))
        trip.status = "departed"
        trip.load_finished_at = now
        trip.departed_at = now
        for stop in list(trip.stops):
            if not stop.boxes:
                db.session.delete(stop)
        for box in trip.boxes:
            if box.status == "loaded":
                _move_box(box, "shipped", now)
        _sync_wms_shipped({box.order for box in trip.boxes}, now)
    elif action == "cancel" and trip.status in TRIP_EDITABLE_STATUSES:
        trip.status = "cancelled"
        for order in list(trip.planned_orders):
            order.planned_trip = None
        for box in list(trip.boxes):
            # погруженные короба возвращаются на склад
            box.status = "received"
            box.loaded_at = None
            box.trip_id = None
            box.trip_stop_id = None
    else:
        flash("Это действие сейчас недоступно", "warning")
        return redirect(url_for("mvb.trip_detail", trip_id=trip.id))
    db.session.commit()
    flash(f"Рейс {trip.number}: {trip.status_label}", "success")
    if _is_driver():
        return redirect(url_for("mvb.driver"))
    return redirect(url_for("mvb.trip_detail", trip_id=trip.id))


@bp.route("/trips/<int:trip_id>/scan", methods=["POST"])
def trip_scan(trip_id):
    """Погрузка сканом (кладовщик): короб или паллета целиком; короб
    попадает на точку маршрута своего направления. Первый скан сам отмечает
    начало погрузки. Сверх вместимости авто — предупреждение."""
    if not _is_staff():
        return jsonify(ok=False, message="Нет доступа"), 403
    trip = _get_trip_or_404(trip_id)
    if trip.status not in ("assigned", "arrived", "loading"):
        return jsonify(ok=False, message=f"Рейс в статусе «{trip.status_label}» — погрузка закрыта"), 409
    code = (request.form.get("barcode") or "").strip()
    now = datetime.utcnow()

    pallet = MvbPallet.query.filter(db.func.upper(MvbPallet.number) == code.upper()).first()
    if pallet is not None:
        boxes = [b for b in pallet.boxes if b.status == "received" and b.trip_id is None]
        if not boxes:
            return jsonify(ok=False, message=f"На паллете {pallet.number} нет коробов к погрузке"), 409
        order = boxes[0].order
    else:
        box = _find_box(code)
        if box is None:
            return jsonify(ok=False, message=f"Короб или паллета {code} не найдены"), 404
        if box.trip_id == trip.id:
            return jsonify(ok=True, already=True, message=f"{box.barcode} уже в этом рейсе", count=len(trip.boxes))
        if box.status != "received" or box.trip_id is not None:
            return jsonify(ok=False, message=f"Короб в статусе «{box.status_label}»"), 409
        boxes = [box]
        order = box.order
    stop = trip.stop_for(order.marketplace, order.destination)
    if stop is None:
        return jsonify(
            ok=False,
            message=f"Направления «{order.marketplace_label} · {order.destination or '—'}» нет в маршруте — добавьте точку",
        ), 409

    if trip.status != "loading":
        trip.status = "loading"
        trip.arrived_at = trip.arrived_at or now
        trip.load_started_at = now
    for box in boxes:
        box.trip = trip
        box.trip_stop = stop
        _move_box(box, "loaded", now)
    db.session.commit()
    count = len(trip.boxes)
    message = (
        f"Паллета {pallet.number}: погружено {len(boxes)} кор. → {stop.label()}" if pallet is not None
        else f"{boxes[0].barcode} погружен → {stop.label()}"
    )
    warnings = []
    capacity = trip.capacity()
    if capacity and count > capacity:
        warnings.append(f"{count} кор. больше вместимости авто ({capacity})")
    if stop.planned_boxes and len(stop.boxes) > stop.planned_boxes:
        warnings.append(f"на точку «{stop.label()}» по плану {stop.planned_boxes} кор., погружено {len(stop.boxes)}")
    if order.planned_trip_id and order.planned_trip_id != trip.id:
        warnings.append(f"заявка {order.number} запланирована в рейс {order.planned_trip.number}")
    warning = ("Внимание: " + "; ".join(warnings)) if warnings else None
    return jsonify(ok=True, already=False, message=message, count=count, warning=warning)


# ---------- ссылка для водителя на СЦ (без входа) ----------
#
# На СЦ чаще всего едут случайные (наемные) водители — их не регистрируем:
# оператор вносит ФИО/телефон/госномер в рейс и отправляет водителю ссылку
# /mvb/t/<токен рейса>, где тот отмечает подачу и итог на каждой точке.


def _trip_by_token_or_404(token):
    trip = MvbTrip.query.filter_by(access_token=token).first() if token else None
    if trip is None or trip.status == "cancelled":
        abort(404)
    return trip


@bp.route("/t/<token>")
def trip_public(token):
    trip = _trip_by_token_or_404(token)
    return render_template("mvb/trip_public.html", trip=trip, token=token)


@bp.route("/t/<token>/arrive", methods=["POST"])
def trip_public_arrive(token):
    trip = _trip_by_token_or_404(token)
    if trip.status == "assigned":
        trip.status = "arrived"
        trip.arrived_at = datetime.utcnow()
        db.session.commit()
        flash("Отмечено: авто на погрузке", "success")
    return redirect(url_for("mvb.trip_public", token=token))


@bp.route("/t/<token>/stops/<int:stop_id>/<action>", methods=["POST"])
def trip_public_stop(token, stop_id, action):
    trip = _trip_by_token_or_404(token)
    stop = db.session.get(MvbTripStop, stop_id)
    if stop is None or stop.trip_id != trip.id or action not in ("deliver", "reject"):
        abort(404)
    category, message = _apply_stop_result(trip, stop, action, request.form.get("comment", ""))
    flash(message, category)
    return redirect(url_for("mvb.trip_public", token=token))


# ---------- свои короба из WMS ----------
#
# Перемещения WMS на склады маркетплейсов (их короба уже с этикетками
# BOX-...) передаются в МВБ как заявки служебного клиента «Свои короба
# (WMS)»; штрихкод короба МВБ = штрихкод короба WMS, поэтому дальше короба
# сканируются по своим же этикеткам, ничего не переклеивая. Когда рейс МВБ
# увозит все короба перемещения, в WMS ставится отметка «Транспорт забрал».

INTERNAL_CLIENT_NAME = "Свои короба (WMS)"
FINISHED_BOX_STATUSES = {"delivered"}


def _internal_client():
    client = MvbClient.query.filter_by(is_internal=True).order_by(MvbClient.id).first()
    if client is None:
        client = MvbClient(name=INTERNAL_CLIENT_NAME, is_internal=True)
        db.session.add(client)
        db.session.flush()
    return client


def _wms_candidates_query():
    """Собранные перемещения WMS на склады WB/Ozon, которые еще не уехали
    (в WMS не отмечено «Транспорт забрал»)."""
    from ..models import MovementDocument, Warehouse

    return (
        MovementDocument.query.join(Warehouse, MovementDocument.to_warehouse_id == Warehouse.id)
        .filter(
            MovementDocument.status == "completed",
            MovementDocument.shipped_at.is_(None),
            Warehouse.marketplace.in_(list(MVB_MARKETPLACES)),
        )
    )


def _active_order_for_movement(doc_id):
    return MvbOrder.query.filter(
        MvbOrder.wms_movement_id == doc_id, MvbOrder.status != "cancelled"
    ).first()


def _find_box(code):
    """Короб МВБ по скану: свой штрихкод МВБ или этикетка короба WMS (в т.ч.
    номер BOX-000123, введенный вручную)."""
    code = (code or "").strip()
    if not code:
        return None
    box = MvbBox.query.filter(db.func.upper(MvbBox.barcode) == code.upper()).first()
    if box is not None:
        return box
    from ..models import Box

    wms_box = Box.find_by_scanned_code(code)
    if wms_box is None:
        return None
    return (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.wms_box_id == wms_box.id, MvbOrder.status != "cancelled")
        .order_by(MvbBox.id.desc())
        .first()
    )


def _wms_movement_for_box_code(code):
    """Перемещение WMS (из кандидатов на передачу), в котором едет короб с
    этой этикеткой и которое еще не передано в МВБ."""
    from ..models import Box, MovementDocument, MovementLine

    wms_box = Box.find_by_scanned_code(code)
    if wms_box is None:
        return None
    docs = (
        _wms_candidates_query()
        .join(MovementLine, MovementLine.document_id == MovementDocument.id)
        .filter(MovementLine.box_id == wms_box.id)
        .all()
    )
    for doc in docs:
        if _active_order_for_movement(doc.id) is None:
            return doc
    return None


def _import_movement(doc, delivery_method="self", pickup_address=None):
    """Создает оформленную заявку МВБ из перемещения WMS. Возвращает
    (заявка или None, список проблем по коробам)."""
    if _active_order_for_movement(doc.id) is not None:
        return None, [f"Перемещение {doc.number} уже передано в МВБ"]
    wms_boxes = []
    seen = set()
    for line in doc.lines:
        if line.box_id not in seen:
            seen.add(line.box_id)
            wms_boxes.append(line.box)
    if not wms_boxes:
        return None, [f"В перемещении {doc.number} нет коробов"]

    errors = []
    accepted = []
    for wms_box in wms_boxes:
        barcode = wms_box.barcode_value
        existing = MvbBox.query.filter(MvbBox.barcode == barcode).first()
        if existing is not None:
            if existing.status in FINISHED_BOX_STATUSES or existing.order.status == "cancelled":
                # Короб WMS используется повторно — прежнюю запись МВБ
                # архивируем, чтобы штрихкод снова был свободен.
                existing.barcode = f"{existing.barcode}#{existing.id}"
            else:
                errors.append(f"Короб {wms_box.box_number} уже в работе МВБ ({existing.order.number})")
                continue
        accepted.append(wms_box)
    if not accepted:
        return None, errors

    destination = doc.to_warehouse
    now = datetime.utcnow()
    order = MvbOrder(
        number=next_number("mvb_order", "MVB-", 6),
        client_id=_internal_client().id,
        marketplace=destination.marketplace,
        destination=destination.marketplace_city or destination.name,
        box_count=len(accepted),
        delivery_method=delivery_method if delivery_method in MVB_DELIVERY_METHODS else "self",
        pickup_address=pickup_address or (doc.from_warehouse.address if doc.from_warehouse else None),
        comment=f"Перемещение WMS {doc.number}: {doc.from_warehouse.name} → {destination.name}",
        status="confirmed",
        confirmed_at=now,
        created_by_id=current_user.id if current_user.is_authenticated else None,
        wms_movement_id=doc.id,
    )
    db.session.add(order)
    db.session.flush()
    for seq, wms_box in enumerate(accepted, start=1):
        order.boxes.append(MvbBox(seq=seq, barcode=wms_box.barcode_value, status="created", wms_box_id=wms_box.id))
    db.session.commit()
    return order, errors


def _sync_wms_shipped(orders, now):
    """Все короба перемещения уехали рейсом МВБ → в WMS «Транспорт забрал»
    (как кнопка на странице перемещения: только если заявка на МП подана)."""
    for order in orders:
        doc = order.wms_movement
        if doc is None or doc.shipped_at is not None or not doc.marketplace_request_created_at:
            continue
        if all(MVB_BOX_STATUS_ORDER.get(b.status, 0) >= MVB_BOX_STATUS_ORDER["shipped"] for b in order.boxes):
            doc.shipped_at = now


@bp.route("/wms")
def wms_movements():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    from ..models import MovementDocument

    docs = _wms_candidates_query().order_by(MovementDocument.completed_at.desc()).limit(300).all()
    rows = []
    for doc in docs:
        rows.append({
            "doc": doc,
            "boxes": len({line.box_id for line in doc.lines}),
            "order": _active_order_for_movement(doc.id),
        })
    imported = (
        MvbOrder.query.filter(MvbOrder.wms_movement_id.isnot(None))
        .order_by(MvbOrder.created_at.desc()).limit(50).all()
    )
    return render_template("mvb/wms.html", rows=rows, imported=imported)


@bp.route("/wms/<int:doc_id>/import", methods=["POST"])
def wms_import(doc_id):
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    from ..models import MovementDocument

    # filter_by здесь применился бы к присоединенной таблице складов.
    doc = _wms_candidates_query().filter(MovementDocument.id == doc_id).first()
    if doc is None:
        flash("Перемещение не найдено или уже уехало", "danger")
        return redirect(url_for("mvb.wms_movements"))
    order, errors = _import_movement(doc, request.form.get("delivery_method", "self"))
    for error in errors:
        flash(error, "warning")
    if order is None:
        return redirect(url_for("mvb.wms_movements"))
    flash(f"Перемещение {doc.number} передано в МВБ: заявка {order.number}, коробов {order.box_count}. Этикетки WMS остаются прежними.", "success")
    return redirect(url_for("mvb.order_detail", order_id=order.id))



# ---------- прайс, стоимость заявки, отчеты ----------


def _apply_prices(order):
    """Стоимость заявки по прайсу: забор — только для способа «забор», для
    своих коробов из WMS стоимость не считается."""
    if order.client and order.client.is_internal:
        return
    order.pickup_cost = (
        MvbPriceTier.cost_for("pickup", order.box_count) if order.delivery_method == "pickup" else None
    )
    order.sc_cost = MvbPriceTier.cost_for("sc", order.box_count)


def _parse_money(value):
    value = (value or "").strip().replace(" ", "").replace(",", ".")
    if not value:
        return None
    try:
        amount = float(value)
    except ValueError:
        raise ValueError(value)
    if amount < 0:
        raise ValueError(value)
    return round(amount, 2)


@bp.route("/orders/<int:order_id>/costs", methods=["POST"])
def order_costs(order_id):
    """Оператор вносит/правит стоимость забора и отправки на СЦ или
    пересчитывает ее по прайсу."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    order = _get_order_or_404(order_id)
    if request.form.get("action") == "recalc":
        _apply_prices(order)
        flash("Стоимость пересчитана по прайсу", "success")
    else:
        try:
            order.pickup_cost = _parse_money(request.form.get("pickup_cost"))
            order.sc_cost = _parse_money(request.form.get("sc_cost"))
        except ValueError:
            flash("Стоимость — неотрицательное число", "danger")
            return redirect(url_for("mvb.order_detail", order_id=order.id))
        flash("Стоимость сохранена", "success")
    db.session.commit()
    return redirect(url_for("mvb.order_detail", order_id=order.id))


@bp.route("/prices", methods=["GET", "POST"])
def prices():
    """Прайс оператора: цена за короб для забора и для отправки на СЦ с
    градацией «от N коробов»."""
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    if request.method == "POST":
        action = request.form.get("action", "save")
        if action == "delete":
            tier = db.session.get(MvbPriceTier, request.form.get("tier_id", type=int)) or abort(404)
            db.session.delete(tier)
            db.session.commit()
            flash("Строка прайса удалена", "success")
            return redirect(url_for("mvb.prices"))
        kind = request.form.get("kind", "")
        min_boxes = request.form.get("min_boxes", type=int)
        try:
            price = _parse_money(request.form.get("price_per_box"))
        except ValueError:
            price = None
        if kind not in MVB_PRICE_KINDS or not min_boxes or min_boxes < 1 or price is None:
            flash("Укажите «от скольких коробов» (от 1) и цену за короб", "danger")
            return redirect(url_for("mvb.prices"))
        tier_id = request.form.get("tier_id", type=int)
        tier = db.session.get(MvbPriceTier, tier_id) if tier_id else None
        duplicate = MvbPriceTier.query.filter_by(kind=kind, min_boxes=min_boxes).first()
        if duplicate is not None and duplicate is not tier:
            # Одна ступень на количество: правим существующую, лишнюю удаляем.
            if tier is not None:
                db.session.delete(tier)
            tier = duplicate
        if tier is None:
            tier = MvbPriceTier(kind=kind)
            db.session.add(tier)
        tier.min_boxes = min_boxes
        tier.price_per_box = price
        db.session.commit()
        flash("Прайс сохранен", "success")
        return redirect(url_for("mvb.prices"))
    tiers = {
        kind: MvbPriceTier.query.filter_by(kind=kind).order_by(MvbPriceTier.min_boxes).all()
        for kind in MVB_PRICE_KINDS
    }
    return render_template("mvb/prices.html", tiers=tiers, kinds=MVB_PRICE_KINDS)


def _report_period():
    today = date.today()
    try:
        date_from = date.fromisoformat(request.args.get("date_from", ""))
    except ValueError:
        date_from = today.replace(day=1)
    try:
        date_to = date.fromisoformat(request.args.get("date_to", ""))
    except ValueError:
        date_to = today
    return date_from, date_to


def _report_rows(date_from, date_to):
    """Отчет по клиентам за период (даты по Москве): заявки, оформленные в
    периоде, их короба по этапам и стоимость; отдельно — сколько коробов
    клиента отправлено на СЦ в периоде (по дате отправки рейса)."""
    from datetime import time

    start = datetime.combine(date_from, time.min) - MOSCOW_OFFSET
    end = datetime.combine(date_to, time.max) - MOSCOW_OFFSET
    rows = {}

    def row_for(client):
        return rows.setdefault(client.id, {
            "client": client, "orders": 0, "boxes": 0, "picked_up": 0, "received": 0,
            "shipped": 0, "delivered": 0, "not_delivered": 0,
            "pickup_cost": 0.0, "sc_cost": 0.0, "total": 0.0,
        })

    orders = MvbOrder.query.filter(
        MvbOrder.status == "confirmed", MvbOrder.confirmed_at >= start, MvbOrder.confirmed_at <= end,
    ).all()
    for order in orders:
        row = row_for(order.client)
        row["orders"] += 1
        row["boxes"] += len(order.boxes)
        row["picked_up"] += sum(1 for b in order.boxes if b.picked_up_at)
        row["received"] += sum(1 for b in order.boxes if b.received_at)
        row["pickup_cost"] += order.pickup_cost or 0
        row["sc_cost"] += order.sc_cost or 0
        row["total"] += order.total_cost or 0

    shipped = (
        MvbBox.query.join(MvbOrder)
        .filter(MvbBox.shipped_at >= start, MvbBox.shipped_at <= end)
        .all()
    )
    for box in shipped:
        row = row_for(box.order.client)
        row["shipped"] += 1
        if box.status == "delivered":
            row["delivered"] += 1
        elif box.status == "not_delivered":
            row["not_delivered"] += 1
    result = sorted(rows.values(), key=lambda r: (-r["shipped"], -r["boxes"], r["client"].name))
    totals = {key: sum(r[key] for r in result) for key in (
        "orders", "boxes", "picked_up", "received", "shipped", "delivered", "not_delivered",
        "pickup_cost", "sc_cost", "total",
    )}
    return result, totals


REPORT_COLUMNS = [
    ("orders", "Заявок"), ("boxes", "Коробов в заявках"), ("picked_up", "Забрано"),
    ("received", "Принято на складе"), ("shipped", "Отправлено на СЦ"), ("delivered", "Сдано на СЦ"),
    ("not_delivered", "Не сдано"), ("pickup_cost", "Забор, руб."), ("sc_cost", "Отправка на СЦ, руб."),
    ("total", "Итого, руб."),
]


@bp.route("/reports")
def reports():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    date_from, date_to = _report_period()
    rows, totals = _report_rows(date_from, date_to)
    return render_template(
        "mvb/reports.html", rows=rows, totals=totals, date_from=date_from, date_to=date_to,
        columns=REPORT_COLUMNS,
    )


@bp.route("/reports.xlsx")
def reports_xlsx():
    if not _require_staff():
        return redirect(url_for("mvb.index"))
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    date_from, date_to = _report_period()
    rows, totals = _report_rows(date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "По клиентам"
    ws.append([f"МВБ Логистика — отчет по клиентам с {date_from:%d.%m.%Y} по {date_to:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Клиент", "ИНН"] + [label for _, label in REPORT_COLUMNS])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row["client"].name, row["client"].inn or ""] + [row[key] for key, _ in REPORT_COLUMNS])
    ws.append(["Итого", ""] + [totals[key] for key, _ in REPORT_COLUMNS])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    for col in "CDEFGHIJKL":
        ws.column_dimensions[col].width = 16
    buffer = io.BytesIO()
    wb.save(buffer)
    fname = f"mvb_report_{date_from:%Y%m%d}_{date_to:%Y%m%d}.xlsx"
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(fname)},
    )
