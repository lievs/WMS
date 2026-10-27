"""Приёмка по акту маркетплейса (movement.receive_act + receive с
act_payload): короба приводятся к принятому количеству, короба вне акта
возвращаются на склад-отправитель, принятое видно в скобках в списке."""

import io
import json
from datetime import datetime

import openpyxl

from wms.extensions import db
from wms.models import (
    Box, BoxItem, MovementCargoPlace, MovementDocument, MovementLine,
    MovementReceiptDiscrepancy, Nomenclature, ShipmentPlan, ShipmentPlanLine,
    UnplacedStock, Warehouse,
)


def _setup(marketplace="wb", ozon=False):
    sender = Warehouse(code="WH-A1", name="Основной")
    dest = Warehouse(code="WH-A2", name="Коледино", marketplace="ozon" if ozon else marketplace)
    db.session.add_all([sender, dest])
    n1 = Nomenclature(sku="A-1", barcode="2000000000011", name="Шапка", unit="шт")
    n2 = Nomenclature(sku="A-2", barcode="2000000000028", barcode2="2000000000099", name="Свитер", unit="шт")
    db.session.add_all([n1, n2])
    db.session.commit()
    plan = ShipmentPlan(marketplace="wb")
    db.session.add(plan)
    db.session.commit()
    pl1 = ShipmentPlanLine(plan_id=plan.id, warehouse_id=dest.id, nomenclature_id=n1.id,
                           barcode=n1.barcode, planned_qty=100, fulfilled_qty=0)
    pl2 = ShipmentPlanLine(plan_id=plan.id, warehouse_id=dest.id, nomenclature_id=n2.id,
                           barcode=n2.barcode, planned_qty=100, fulfilled_qty=0)
    db.session.add_all([pl1, pl2])
    boxes = []
    for number, items in (("BOX-A1", {n1: 10}), ("BOX-A2", {n2: 5}), ("BOX-A3", {n1: 4, n2: 6})):
        box = Box(box_number=number, warehouse_id=dest.id, status="open")
        db.session.add(box)
        db.session.commit()
        for nom, qty in items.items():
            db.session.add(BoxItem(box_id=box.id, nomenclature_id=nom.id, qty=qty))
        boxes.append(box)
    doc = MovementDocument(
        number="PER-A0001", from_warehouse_id=sender.id, to_warehouse_id=dest.id, status="completed",
        marketplace_request_created_at=datetime.utcnow(), marketplace_request_number="REQ-1",
        shipped_at=datetime.utcnow(),
    )
    db.session.add(doc)
    db.session.commit()
    for box in boxes:
        db.session.add(MovementLine(document_id=doc.id, box_id=box.id, from_warehouse_id=sender.id))
    db.session.commit()
    return doc, sender, dest, n1, n2, boxes, (pl1, pl2)


def _act_file(rows):
    """rows: (баркод, ШК короба, кол-во) — формат акта WB: шапка в 2 строки,
    строка нумерации колонок, итог."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None])
    ws.append(["Акт приемки товара № 1"])
    ws.append([])
    ws.append(["№ п\\п", "Товар (наименование)", "Ед. изм.", "Фактически принято"])
    ws.append([None, None, None, "баркод", "артикул продавца", "сорт, размер", "КИЗ", "ШК короба", "кол-во"])
    ws.append([None, "1", "2", "3", "4", "5", "6", "7", "8"])
    for i, (barcode, box, qty) in enumerate(rows, 1):
        ws.append([i, "Товар", "шт.", barcode, "арт", "50", f"kiz{i}", box, qty])
    ws.append(["Итого", None, None, None, None, None, None, None, sum(r[2] for r in rows)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _upload(client, doc, rows):
    return client.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (_act_file(rows), "act.xlsx")},
        content_type="multipart/form-data",
    )


def _confirm(client, doc, html):
    import re
    from html import unescape

    payload = unescape(re.search(r'name="act_payload" value="([^"]*)"', html).group(1))
    return client.post(f"/movement/{doc.id}/receive", data={"act_payload": payload})


def test_act_preview_shows_diff_and_missing_boxes_without_writing(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    resp = _upload(client_logged_in, doc, [
        (n1.barcode, "BOX-A1", 1), ("2000000000011", "BOX-A1", 1),  # 2 из 10
        ("02000000000099", "BOX-A3", 6),  # доп. штрихкод с ведущим нулём
    ])
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200
    assert "Не приняты вовсе" in html and "BOX-A2" in html
    assert "act_payload" in html
    assert MovementDocument.query.get(doc.id).received_at is None
    assert BoxItem.query.filter_by(box_id=boxes[0].id).first().qty == 10


def test_act_confirm_trims_boxes_returns_missing_and_records_received(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, (pl1, pl2) = _setup()
    html = _upload(client_logged_in, doc, [
        (n1.barcode, "BOX-A1", 8),
        (n2.barcode, "BOX-A3", 6),  # из BOX-A3 приняли только свитера, шапок нет
    ]).get_data(as_text=True)
    resp = _confirm(client_logged_in, doc, html)
    assert resp.status_code == 302

    doc = MovementDocument.query.get(doc.id)
    assert doc.received_at is not None and doc.receipt_from_act is True
    assert doc.receipt_returned_boxes == "BOX-A2"
    # BOX-A2 возвращён отправителю и снят с перемещения
    assert Box.query.get(boxes[1].id).warehouse_id == sender.id
    assert doc.lines.count() == 2
    # короба приведены к акту
    assert BoxItem.query.filter_by(box_id=boxes[0].id).first().qty == 8
    a3 = {i.nomenclature_id: i.qty for i in BoxItem.query.filter_by(box_id=boxes[2].id)}
    assert a3 == {n2.id: 6}
    # отправлено 25, принято 14 → в скобках
    assert doc.total_sent_qty() == 25 and doc.total_received_qty() == 14
    # план: засчитано только принятое
    assert ShipmentPlanLine.query.get(pl1.id).fulfilled_qty == 8
    assert ShipmentPlanLine.query.get(pl2.id).fulfilled_qty == 6
    d = {x.nomenclature_id: (x.expected_qty, x.received_qty) for x in MovementReceiptDiscrepancy.query}
    assert d == {n1.id: (14, 8)}


def test_act_excess_in_box_debits_sender_unplaced_stock(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    UnplacedStock.add(sender.id, n1.id, 5)
    db.session.commit()
    html = _upload(client_logged_in, doc, [
        (n1.barcode, "BOX-A1", 12), (n2.barcode, "BOX-A2", 5),
        (n1.barcode, "BOX-A3", 4), (n2.barcode, "BOX-A3", 6),
    ]).get_data(as_text=True)
    _confirm(client_logged_in, doc, html)
    assert BoxItem.query.filter_by(box_id=boxes[0].id).first().qty == 12
    assert UnplacedStock.available(sender.id, n1.id) == 3
    assert MovementDocument.query.get(doc.id).receipt_returned_boxes is None


def test_act_unknown_barcode_blocks_confirmation(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    html = _upload(client_logged_in, doc, [("9999999999999", "BOX-A1", 3)]).get_data(as_text=True)
    assert "нет в номенклатуре" in html
    assert "act_payload" not in html


def test_act_with_no_matching_boxes_is_rejected(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    html = _upload(client_logged_in, doc, [(n1.barcode, "BOX-ZZ", 3)]).get_data(as_text=True)
    assert "нечего принимать" in html
    assert "act_payload" not in html


def test_act_bad_file_shows_message(db, client_logged_in):
    doc, *_ = _setup()
    resp = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (io.BytesIO(b"not an excel"), "x.xlsx")},
        content_type="multipart/form-data", follow_redirects=True,
    )
    assert "Не удалось прочитать файл" in resp.get_data(as_text=True)


def test_unreceive_after_act_restores_trimmed_quantity(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, (pl1, pl2) = _setup()
    html = _upload(client_logged_in, doc, [
        (n1.barcode, "BOX-A1", 8), (n2.barcode, "BOX-A2", 5),
        (n1.barcode, "BOX-A3", 4), (n2.barcode, "BOX-A3", 6),
    ]).get_data(as_text=True)
    _confirm(client_logged_in, doc, html)
    assert ShipmentPlanLine.query.get(pl1.id).fulfilled_qty == 12
    client_logged_in.post(f"/movement/{doc.id}/unreceive")
    doc = MovementDocument.query.get(doc.id)
    assert doc.received_at is None and doc.receipt_from_act is None
    assert ShipmentPlanLine.query.get(pl1.id).fulfilled_qty == 0
    total_n1 = sum(i.qty for i in BoxItem.query.filter_by(nomenclature_id=n1.id))
    assert total_n1 == 14


def test_ozon_act_matches_boxes_by_remembered_cargo_places(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup(ozon=True)
    # без запомненного сопоставления — понятная ошибка
    html = _upload(client_logged_in, doc, [(n1.barcode, "GM-1", 10)]).get_data(as_text=True)
    assert "выгрузите «Состав ГМ" in html

    db.session.add_all([
        MovementCargoPlace(document_id=doc.id, box_id=boxes[0].id, box_number="BOX-A1", cargo_barcode="GM-1"),
        MovementCargoPlace(document_id=doc.id, box_id=boxes[1].id, box_number="BOX-A2", cargo_barcode="GM-2"),
        MovementCargoPlace(document_id=doc.id, box_id=boxes[2].id, box_number="BOX-A3", cargo_barcode="GM-3"),
    ])
    db.session.commit()
    html = _upload(client_logged_in, doc, [(n1.barcode, "GM-1", 9), (n2.barcode, "GM-2", 5)]).get_data(as_text=True)
    assert "BOX-A3" in html and "Не приняты вовсе" in html
    _confirm(client_logged_in, doc, html)
    doc = MovementDocument.query.get(doc.id)
    assert doc.receipt_returned_boxes == "BOX-A3"
    assert BoxItem.query.filter_by(box_id=boxes[0].id).first().qty == 9
    assert MovementCargoPlace.query.filter_by(document_id=doc.id).count() == 2


def test_act_cannot_be_applied_twice(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    rows = [(n1.barcode, "BOX-A1", 10), (n2.barcode, "BOX-A2", 5), (n1.barcode, "BOX-A3", 4), (n2.barcode, "BOX-A3", 6)]
    html = _upload(client_logged_in, doc, rows).get_data(as_text=True)
    _confirm(client_logged_in, doc, html)
    resp = _confirm(client_logged_in, doc, html)
    assert resp.status_code == 302
    assert MovementReceiptDiscrepancy.query.count() == 0
    assert ShipmentPlanLine.query.filter_by(nomenclature_id=n1.id).first().fulfilled_qty == 14


def test_act_zip_archive_is_accepted_as_is(db, client_logged_in):
    import zipfile

    doc, sender, dest, n1, n2, boxes, _ = _setup()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("signature.sig", b"sig")
        zf.writestr("act-income-1.xlsx", _act_file([(n1.barcode, "BOX-A1", 3)]).getvalue())
        zf.writestr("mchd.zip", b"PK-not-an-act")
    buf.seek(0)
    resp = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (buf, "Acceptance-certificate.zip")},
        content_type="multipart/form-data",
    )
    html = resp.get_data(as_text=True)
    assert "act_payload" in html and "BOX-A1" in html


def test_zip_without_xlsx_is_rejected(db, client_logged_in):
    import zipfile

    doc, *_ = _setup()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", b"x")
    buf.seek(0)
    resp = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (buf, "a.zip")}, content_type="multipart/form-data", follow_redirects=True,
    )
    assert "нет файла акта" in resp.get_data(as_text=True)


def test_fast_and_openpyxl_readers_give_same_result():
    from wms.utils import acceptance_act as aa

    data = _act_file([("2000000000011", "BOX-A1", 2), ("2000000000028", "BOX-A2", 3), ("2000000000011", "BOX-A1", 1)]).getvalue()
    fast = aa.parse_act(data)
    original = aa._iter_sheet_rows_fast
    try:
        def boom(d, w):
            raise RuntimeError("force fallback")
        aa._iter_sheet_rows_fast = boom
        slow = aa.parse_act(data)
    finally:
        aa._iter_sheet_rows_fast = original
    assert fast == slow
    assert dict(fast[0]) == {"BOX-A1": {"2000000000011": 3.0}, "BOX-A2": {"2000000000028": 3.0}}


def test_receive_pages_show_barcode_column(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    html = client_logged_in.get(f"/movement/{doc.id}/receive").get_data(as_text=True)
    assert "Штрихкод" in html and n1.barcode in html and n2.barcode in html and n2.barcode2 in html
    html = _upload(client_logged_in, doc, [(n1.barcode, "BOX-A1", 3)]).get_data(as_text=True)
    assert "Штрихкод" in html and n1.barcode in html


def _ozon_report(rows):
    """Отчёт Ozon «supply-…-acceptance-report.xlsx»: лист по поставке (без
    ГМ) и лист «Отчёт по ГМ» (ШК ГМ, ШК товара, заявлено, принято)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт по поставке"
    ws.append(["Принято ГМ / Заявлено ГМ", "3/3"])
    ws.append([])
    ws.append(["Ozon ID", "ШК", "Артикул товара", "Наименование товара", "Заявлено (шт.)", "Принято на склад (шт.)"])
    ws2 = wb.create_sheet("Отчёт по ГМ")
    ws2.append(["Принято ГМ / Заявлено ГМ", "3/3"])
    ws2.append([])
    ws2.append(["ШК ГМ", "ГМ", "Тип ГМ", "Ozon ID", "ШК товара", "Артикул товара", "Наименование товара",
                "Заявлено (шт.)", "Принято на склад (шт.)", "Разница по ГМ"])
    for gm, barcode, declared, accepted, *article in rows:
        ws2.append([gm, "Коробка", "Моно", 1, barcode, article[0] if article else "арт", "Товар",
                    declared, accepted, accepted - declared])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_ozon_acceptance_report_layout_is_recognized(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup(ozon=True)
    db.session.add_all([
        MovementCargoPlace(document_id=doc.id, box_id=boxes[0].id, box_number="BOX-A1", cargo_barcode="1000000000000001"),
        MovementCargoPlace(document_id=doc.id, box_id=boxes[1].id, box_number="BOX-A2", cargo_barcode="1000000000000002"),
        MovementCargoPlace(document_id=doc.id, box_id=boxes[2].id, box_number="BOX-A3", cargo_barcode="1000000000000003"),
    ])
    db.session.commit()
    report = _ozon_report([
        ("1000000000000001", "0" + n1.barcode, 10, 7),   # ведущий ноль, недостача
        ("1000000000000002", n2.barcode, 5, 5),
        ("1000000000000003", n1.barcode, 4, 4),
        ("1000000000000003", n2.barcode, 6, 6),
    ])
    resp = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (report, "supply-1-acceptance-report.xlsx")}, content_type="multipart/form-data",
    )
    html = resp.get_data(as_text=True)
    assert "act_payload" in html and "Не приняты вовсе" not in html
    _confirm(client_logged_in, doc, html)
    assert BoxItem.query.filter_by(box_id=boxes[0].id).first().qty == 7
    assert MovementDocument.query.get(doc.id).total_received_qty() == 22


def test_ozon_act_matches_goods_by_ozon_article_not_barcode(db, client_logged_in):
    from wms.models import OzonArticleMapping

    doc, sender, dest, n1, n2, boxes, _ = _setup(ozon=True)
    db.session.add_all([
        MovementCargoPlace(document_id=doc.id, box_id=boxes[0].id, box_number="BOX-A1", cargo_barcode="1000000000000001"),
        MovementCargoPlace(document_id=doc.id, box_id=boxes[1].id, box_number="BOX-A2", cargo_barcode="1000000000000002"),
        MovementCargoPlace(document_id=doc.id, box_id=boxes[2].id, box_number="BOX-A3", cargo_barcode="1000000000000003"),
        OzonArticleMapping(barcode=n1.barcode, article="ШАПКА-1"),
        OzonArticleMapping(barcode=n2.barcode, article="СВИТЕР-2"),
    ])
    db.session.commit()
    report = _ozon_report([
        # баркод в отчёте искажён (потерян/лишний знак), товар определяется по артикулу
        ("1000000000000001", "999", 10, 7, "ШАПКА-1"),
        ("1000000000000002", "", 5, 5, "свитер-2"),  # регистр артикула не важен
        # артикула нет в справочнике — запасной путь по баркоду строки
        ("1000000000000003", "0" + n1.barcode, 4, 4, "НЕТ-В-СПРАВОЧНИКЕ"),
        ("1000000000000003", n2.barcode, 6, 6, "СВИТЕР-2"),
    ])
    resp = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (report, "r.xlsx")}, content_type="multipart/form-data",
    )
    html = resp.get_data(as_text=True)
    assert "act_payload" in html
    _confirm(client_logged_in, doc, html)
    assert BoxItem.query.filter_by(box_id=boxes[0].id, nomenclature_id=n1.id).first().qty == 7
    assert MovementDocument.query.get(doc.id).total_received_qty() == 22


def test_ozon_act_unmapped_article_and_unknown_barcode_blocks(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup(ozon=True)
    db.session.add(MovementCargoPlace(document_id=doc.id, box_id=boxes[0].id, box_number="BOX-A1", cargo_barcode="1000000000000001"))
    db.session.commit()
    report = _ozon_report([("1000000000000001", "123", 10, 7, "НЕТ-ТАКОГО")])
    html = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (report, "r.xlsx")}, content_type="multipart/form-data",
    ).get_data(as_text=True)
    assert "НЕТ-ТАКОГО" in html and "act_payload" not in html


def test_act_unidentified_row_is_counted_separately_and_totals_shown(db, client_logged_in):
    doc, sender, dest, n1, n2, boxes, _ = _setup()
    rows = [(n1.barcode, "BOX-A1", 10), ("", "BOX-A1", 1), (n2.barcode, "BOX-A2", 5)]
    html = client_logged_in.post(
        f"/movement/{doc.id}/receive/act",
        data={"act_file": (_act_file(rows), "a.xlsx")}, content_type="multipart/form-data",
    ).get_data(as_text=True)
    assert "Неопознан" in html and "В акте всего" in html
    assert "act_payload" in html  # не блокирует подтверждение
    import re
    from html import unescape
    payload = json.loads(unescape(re.search(r'name="act_payload" value="([^"]*)"', html).group(1)))
    assert payload == {"BOX-A1": {str(n1.id): 10.0}, "BOX-A2": {str(n2.id): 5.0}}
