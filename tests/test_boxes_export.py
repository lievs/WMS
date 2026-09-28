"""Выгрузка непустых коробов, разбитых по складам (см. чат) — отдельный
лист Excel на каждый склад, пустые короба (без единой позиции) в выгрузку
не попадают вовсе."""

from io import BytesIO

from openpyxl import load_workbook

from wms.extensions import db
from wms.models import Box, BoxItem, Nomenclature, Warehouse


def _make_item(suffix):
    item = Nomenclature(sku=f"SKU-BXEXP{suffix}", barcode=f"77709000{suffix}", name="Товар", unit="шт")
    db.session.add(item)
    db.session.commit()
    return item


def test_export_only_includes_non_empty_boxes(db, client_logged_in):
    wh = Warehouse(code="WH-BXEXP1", name="Склад А")
    db.session.add(wh)
    db.session.commit()
    item = _make_item("1")

    full_box = Box(box_number="BOX-BXEXP-FULL", warehouse_id=wh.id, status="open")
    empty_box = Box(box_number="BOX-BXEXP-EMPTY", warehouse_id=wh.id, status="open")
    db.session.add_all([full_box, empty_box])
    db.session.commit()
    db.session.add(BoxItem(box_id=full_box.id, nomenclature_id=item.id, qty=5))
    db.session.commit()

    resp = client_logged_in.get("/boxes/export.xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(BytesIO(resp.data))
    ws = wb["Склад А"]
    box_numbers = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value]
    assert box_numbers == ["BOX-BXEXP-FULL"]


def test_export_splits_boxes_into_one_sheet_per_warehouse(db, client_logged_in):
    wh1 = Warehouse(code="WH-BXEXP2", name="Основной склад")
    wh2 = Warehouse(code="WH-BXEXP3", name="Склад №2 (Шоссейная 167)")
    db.session.add_all([wh1, wh2])
    db.session.commit()
    item = _make_item("2")

    box1 = Box(box_number="BOX-BXEXP-W1", warehouse_id=wh1.id, status="open")
    box2 = Box(box_number="BOX-BXEXP-W2", warehouse_id=wh2.id, status="open")
    db.session.add_all([box1, box2])
    db.session.commit()
    db.session.add(BoxItem(box_id=box1.id, nomenclature_id=item.id, qty=1))
    db.session.add(BoxItem(box_id=box2.id, nomenclature_id=item.id, qty=1))
    db.session.commit()

    resp = client_logged_in.get("/boxes/export.xlsx")
    wb = load_workbook(BytesIO(resp.data))

    assert "Основной склад" in wb.sheetnames
    assert "Склад №2 (Шоссейная 167)" in wb.sheetnames
    assert [c.value for c in next(wb["Основной склад"].iter_rows(min_row=2))][0] == "BOX-BXEXP-W1"
    assert [c.value for c in next(wb["Склад №2 (Шоссейная 167)"].iter_rows(min_row=2))][0] == "BOX-BXEXP-W2"


def test_export_respects_warehouse_filter(db, client_logged_in):
    wh1 = Warehouse(code="WH-BXEXP4", name="Склад Раз")
    wh2 = Warehouse(code="WH-BXEXP5", name="Склад Два")
    db.session.add_all([wh1, wh2])
    db.session.commit()
    item = _make_item("3")

    box1 = Box(box_number="BOX-BXEXP-F1", warehouse_id=wh1.id, status="open")
    box2 = Box(box_number="BOX-BXEXP-F2", warehouse_id=wh2.id, status="open")
    db.session.add_all([box1, box2])
    db.session.commit()
    db.session.add(BoxItem(box_id=box1.id, nomenclature_id=item.id, qty=1))
    db.session.add(BoxItem(box_id=box2.id, nomenclature_id=item.id, qty=1))
    db.session.commit()

    resp = client_logged_in.get(f"/boxes/export.xlsx?warehouse_id={wh1.id}")
    wb = load_workbook(BytesIO(resp.data))

    assert wb.sheetnames == ["Склад Раз"]


def test_export_with_no_boxes_returns_empty_workbook(db, client_logged_in):
    resp = client_logged_in.get("/boxes/export.xlsx")

    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Короба"]


def test_boxes_list_has_export_button(db, client_logged_in):
    html = client_logged_in.get("/boxes/").get_data(as_text=True)
    assert "/boxes/export.xlsx" in html
